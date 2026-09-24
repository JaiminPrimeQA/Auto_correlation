from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(r"D:\Auto_correlation")
OUT = ROOT / "reports" / "Baseline11_Auto_Correlation_Tester_Guide.docx"
SCREENSHOTS = [
    Path(r"C:\Users\Almas\AppData\Local\Temp\codex-clipboard-89d30d80-0bc6-46e5-85fe-1e4669c390cc.png"),
    Path(r"C:\Users\Almas\AppData\Local\Temp\codex-clipboard-88533f9a-e708-490d-94fd-9deb54fab61d.png"),
    Path(r"C:\Users\Almas\AppData\Local\Temp\codex-clipboard-54d5a5f1-011a-4ca9-ae22-610273dc36d2.png"),
]
ASSET_DIR = ROOT / "reports" / "assets"
GENERATE_CROP = ASSET_DIR / "generate-and-validate.png"

BLUE = "17365D"
LIGHT_BLUE = "EAF2F8"
PALE = "F7F9FB"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(85, 85, 85)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), GRID)


def set_col_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        shade_cell(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            set_col_width(cell, widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                set_col_width(cell, widths[idx])
            if ridx % 2:
                shade_cell(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(text))
            run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.05
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    p_pr.append(shd)
    for idx, line in enumerate(text.splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")
        run.font.size = Pt(8.4)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(9)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


doc = Document()
section = doc.sections[0]
doc.settings.odd_and_even_pages_header_footer = False
section.different_first_page_header_footer = False
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.12

for name, size, before, after in (
    ("Title", 26, 0, 16),
    ("Heading 1", 17, 14, 7),
    ("Heading 2", 12.5, 10, 5),
    ("Heading 3", 10.5, 8, 3),
):
    style = styles[name]
    style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
    style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
    style.font.color.rgb = BLACK
    style.font.size = Pt(size)
    style.font.bold = name != "Title" or True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Some Word installations add a theme border to the built-in Title style.
# Remove it so the title is separated by whitespace only.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

footer = section.footer.paragraphs[0]
footer.style = styles["Normal"]
footer.runs.clear()
add_page_number(footer)

ASSET_DIR.mkdir(parents=True, exist_ok=True)
if SCREENSHOTS[2].exists():
    with Image.open(SCREENSHOTS[2]) as source:
        width, height = source.size
        source.crop((50, min(430, height - 1), width - 50, min(1110, height))).save(GENERATE_CROP)

# Title page
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("Baseline11 Auto Correlation Tester Guide")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(22)
r = subtitle.add_run("From Postman collection execution to a validated JMeter 5.6.3 test plan")
r.font.size = Pt(15)
r.font.color.rgb = GRAY

intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(14)
intro.add_run("Purpose. ").bold = True
intro.add_run(
    "This guide explains how a tester uses Baseline11 Auto Correlate to execute a working Postman collection twice with Newman, upload the resulting JSON reports, confirm dynamic producer and consumer relationships, generate a correlated JMX, and validate it with Apache JMeter."
)

conclusion = doc.add_paragraph()
conclusion.paragraph_format.space_after = Pt(18)
conclusion.add_run("Expected outcome. ").bold = True
conclusion.add_run(
    "The tester receives a structurally valid JMX in which dynamic response values are extracted and reused automatically, while credentials remain external JMeter properties unless the tester explicitly chooses to embed them."
)

add_table(
    doc,
    ["Document", "Audience", "Version", "Date"],
    [["Operating guide", "API, QA and performance testers", "1.0", "23 September 2026"]],
    [1.35, 2.45, 0.75, 1.55],
)

doc.add_paragraph("What the tester needs", style="Heading 2")
for item in [
    "A Postman collection that already executes successfully.",
    "A Postman environment or another safe source for host and credential values.",
    "Newman and the optional htmlextra reporter.",
    "Two successful Newman JSON reports from the same collection and request order.",
    "Access to the Baseline11 Auto Correlate web application and Apache JMeter 5.6.3.",
]:
    add_bullet(doc, item)

doc.add_page_break()

# Quick start
doc.add_heading("Quick Start", level=1)
doc.add_paragraph(
    "Use this page when the collection, environment and credentials are already prepared. The detailed procedure begins on the next page."
)
quick_rows = [
    ("1", "Verify the collection", "Run the selected folder in Postman or Newman and confirm that producer requests succeed."),
    ("2", "Create baseline", "Run Newman and export baseline.json."),
    ("3", "Create comparison", "Run the same command again and export comparison.json."),
    ("4", "Upload", "Upload baseline first and comparison second."),
    ("5", "Review", "Confirm healthy runs, aligned requests and the expected candidate."),
    ("6", "Correlate", "Click Auto correlate all reused values once."),
    ("7", "Generate", "Preview, generate and download the JMX and manifest."),
    ("8", "Validate", "Supply runtime secrets and execute the plan with JMeter 5.6.3."),
]
add_table(doc, ["Step", "Action", "Expected result"], quick_rows, [0.55, 1.55, 4.05])

doc.add_heading("The Core Concept", level=2)
doc.add_paragraph(
    "Newman does not convert a Postman collection into JSON. Newman executes the collection and records the resolved requests, responses, status codes, assertions and timings in a JSON report. Baseline11 compares two reports to prove which response values are dynamic and where later requests reuse them."
)
add_code(
    doc,
    "Postman collection + environment\n"
    "        |\n"
    "        +--> Newman baseline run   --> baseline.json\n"
    "        +--> Newman comparison run --> comparison.json\n"
    "                                      |\n"
    "                                      +--> Baseline11 --> correlated JMX --> JMeter",
)

doc.add_heading("Why Two Runs", level=2)
doc.add_paragraph(
    "A single run can show that a value is reused, but it cannot reliably distinguish a dynamic identifier from static configuration. Two successful runs show that the value changes while the producer response path and consumer request locations remain stable."
)

doc.add_page_break()

# Preparation
doc.add_heading("1 Prepare the Postman Collection", level=1)
doc.add_paragraph(
    "The collection must complete the intended API journey before it is uploaded indirectly through Newman reports. Baseline11 translates an observed working flow into JMeter; it does not repair authentication failures or unresolved Postman placeholders."
)

doc.add_heading("Collection readiness checklist", level=2)
for item in [
    "Requests are ordered in the same sequence the business flow requires.",
    "Authentication resolves from the selected environment.",
    "Producer requests return successful status codes and usable response bodies.",
    "Postman scripts capture values required for the collection itself to continue.",
    "Later requests use Postman variables such as {{order_id}} instead of literal {order_id} text.",
    "Assertions fail when a required producer response or identifier is missing.",
]:
    add_bullet(doc, item)

doc.add_heading("Example dynamic flow", level=2)
doc.add_paragraph("Create Order returns:")
add_code(doc, '{\n  "id": "order_ABC123",\n  "amount": 10000,\n  "status": "created"\n}')
doc.add_paragraph("A Postman post-response script keeps the collection executable:")
add_code(
    doc,
    'const response = pm.response.json();\n'
    'pm.test("Order ID was returned", function () {\n'
    '    pm.expect(response.id).to.be.a("string").and.not.empty;\n'
    '});\n'
    'pm.collectionVariables.set("order_id", response.id);',
)
doc.add_paragraph("A downstream Postman request uses:")
add_code(doc, "GET https://api.example.com/v1/orders/{{order_id}}")

doc.add_heading("Environment guidance", level=2)
doc.add_paragraph(
    "Keep real keys and secrets in a local environment file. Share only a template with empty values. Confirm that exported environment values are populated before running Newman; some Postman workflows export variable definitions without the local values."
)

doc.add_page_break()

# Newman
doc.add_heading("2 Generate Newman Reports", level=1)
doc.add_heading("Install the command line tools", level=2)
add_code(doc, "npm install -g newman newman-reporter-htmlextra\nnewman --version")

doc.add_heading("Generic baseline command", level=2)
add_code(
    doc,
    'newman run "C:\\QA\\collection.postman_collection.json" ^\n'
    '  -e "C:\\QA\\test.postman_environment.json" ^\n'
    '  --folder "Business Flow" ^\n'
    '  -r cli,json,htmlextra ^\n'
    '  --reporter-json-export "C:\\QA\\reports\\baseline.json" ^\n'
    '  --reporter-htmlextra-export "C:\\QA\\reports\\baseline.html"',
)
doc.add_paragraph("Run the same command again with comparison output names:")
add_code(
    doc,
    'newman run "C:\\QA\\collection.postman_collection.json" ^\n'
    '  -e "C:\\QA\\test.postman_environment.json" ^\n'
    '  --folder "Business Flow" ^\n'
    '  -r cli,json,htmlextra ^\n'
    '  --reporter-json-export "C:\\QA\\reports\\comparison.json" ^\n'
    '  --reporter-htmlextra-export "C:\\QA\\reports\\comparison.html"',
)

doc.add_heading("Command reference", level=2)
add_table(
    doc,
    ["Argument", "Purpose"],
    [
        ("newman run", "Executes a Postman collection."),
        ("-e", "Loads a Postman environment."),
        ("--folder", "Limits execution to one folder or scenario."),
        ("-r cli,json,htmlextra", "Prints terminal output and creates JSON and HTML reports."),
        ("--reporter-json-export", "Writes the machine-readable report consumed by Baseline11."),
        ("--reporter-htmlextra-export", "Writes a human-readable execution report for review."),
    ],
    [2.05, 4.1],
)

doc.add_heading("Report acceptance criteria", level=2)
for item in [
    "Both runs use the same collection, folder and request order.",
    "The producer and intended consumers return successful responses.",
    "There are no authentication, network or missing-response failures.",
    "Required assertions pass.",
    "Dynamic values differ between the two runs.",
]:
    add_bullet(doc, item)

doc.add_page_break()

# Upload and review
doc.add_heading("3 Upload and Review the Analysis", level=1)
doc.add_paragraph(
    "Open Baseline11 and upload baseline.json first and comparison.json second. Do not upload the Postman collection, environment, HTML report or a JMX into the report uploader."
)

doc.add_heading("Run Health", level=2)
doc.add_paragraph(
    "Run Health verifies transport success, assertion results, business response heuristics and sequence alignment. A warning is reviewable; a blocker means the reports are not safe for automatic correlation."
)
if SCREENSHOTS[0].exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(SCREENSHOTS[0]), width=Inches(6.45))
    picture._inline.docPr.set("descr", "Baseline11 Run Health showing seven aligned successful requests and one reviewable empty dataset warning")
    add_caption(doc, "Figure 1  Example Run Health with seven aligned HTTP 200 requests")

doc.add_heading("How to interpret common health results", level=2)
add_table(
    doc,
    ["Result", "Meaning", "Tester action"],
    [
        ("Ready", "Both reports are suitable for analysis.", "Continue to Candidates."),
        ("Ready with review", "Transport is usable but a business heuristic needs review.", "Read the warning and continue only when the response is valid for the scenario."),
        ("Blocker", "Authentication, response availability or alignment is too poor.", "Fix the collection or credentials and create two new reports."),
        ("Empty dataset", "A successful response returned no records.", "Confirm whether an empty result is expected, such as payments for a new unpaid order."),
    ],
    [1.25, 2.45, 2.45],
)

doc.add_page_break()

doc.add_heading("4 Review and Accept Correlation", level=1)
doc.add_paragraph(
    "Candidates are response values that changed between runs and were reused by later requests in both runs. Inspect the producer path, consumer count, confidence and evidence before acceptance."
)
if SCREENSHOTS[1].exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(SCREENSHOTS[1]), width=Inches(6.45))
    picture._inline.docPr.set("descr", "Baseline11 Candidates page showing one high confidence ID value produced at JSONPath id and reused by three requests")
    add_caption(doc, "Figure 2  One high-confidence ID candidate produced at JSONPath $.id and reused by three requests")

doc.add_heading("Decision rules", level=2)
add_table(
    doc,
    ["Observation", "Decision"],
    [
        ("Earlier response produces the value and later requests reuse it", "Accept or use automatic correlation."),
        ("Value changes but no earlier response produces it", "Treat it as input parameterization or runtime configuration."),
        ("Authorization, API key or secret", "Keep it as an external JMeter property."),
        ("Cookie maintained by session behavior", "Use JMeter HTTP Cookie Manager."),
        ("Host, Content-Length, Postman-Token or Accept-Encoding", "Treat as transport noise and let JMeter compute it."),
    ],
    [3.15, 3.0],
)

doc.add_heading("Automatic option", level=2)
add_number(doc, "Click Auto correlate all reused values once.")
add_number(doc, "Confirm the button changes to Correlation completed and becomes disabled.")
add_number(doc, "Open Dependency graph and verify the producer points to the expected consumers.")
add_number(doc, "Confirm Generate displays the accepted rule count.")

doc.add_paragraph(
    "Variable names may be edited for clarity. For example, a response field named id can be renamed from ${id} to ${order_id} when the business meaning is an order identifier."
)

doc.add_page_break()

# Generate
doc.add_heading("5 Generate the JMeter Plan", level=1)
doc.add_paragraph(
    "Open Generate after accepting the correlation. Use Preview Draft to inspect the planned producer, extractor, variable and consumers before creating the JMX."
)

doc.add_heading("Recommended generation settings", level=2)
add_table(
    doc,
    ["Setting", "Recommended value", "Reason"],
    [
        ("Threads", "1 for functional validation", "Reduces diagnostic noise before load sizing."),
        ("Loops", "1 for functional validation", "Confirms one complete business journey."),
        ("Parameterize host", "Enabled", "Keeps the target environment configurable."),
        ("Cache Manager", "Enabled when browser-like caching is desired", "Replays common HTTP caching behavior."),
        ("Keep User-Agent", "Disabled unless required", "Lets JMeter manage transport headers."),
        ("Embed static secrets", "Disabled", "Prevents credentials from being stored in the JMX."),
    ],
    [1.55, 2.05, 2.55],
)

if GENERATE_CROP.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = p.add_run().add_picture(str(GENERATE_CROP), width=Inches(5.7))
    picture._inline.docPr.set("descr", "Baseline11 Generate page showing one accepted correlation, three downstream substitutions, JMX download, and Authorization validation input")
    add_caption(doc, "Figure 3  Generated JMX summary and runtime Authorization property required for validation")

doc.add_heading("What the JMX should contain", level=2)
for item in [
    "One HTTP sampler per aligned Newman execution.",
    "A JSON Extractor attached to the producer request, for example $.id.",
    "JMeter variable references such as ${id} or ${order_id} in downstream requests.",
    "Cookie and cache managers where selected.",
    "External property expressions such as ${__P(Authorization,)} for credentials.",
]:
    add_bullet(doc, item)

doc.add_page_break()

# Validation
doc.add_heading("6 Validate with Apache JMeter 5.6.3", level=1)
doc.add_paragraph(
    "Generated means the XML structure is valid. Validated means Apache JMeter executed the plan successfully with the required runtime properties. The web validation form passes secrets only to that execution and does not store them in the JMX."
)

doc.add_heading("Authorization value for Basic authentication", level=2)
doc.add_paragraph(
    "Supply the complete Basic Authorization header value, not only the key or secret. It has this form:"
)
add_code(doc, "Basic <Base64(api_key:api_secret)>")
doc.add_paragraph("Create the value locally without printing the secret:")
add_code(
    doc,
    "$e = Get-Content -Raw -LiteralPath 'C:\\QA\\test.postman_environment.json' | ConvertFrom-Json\n"
    "$key = ($e.values | Where-Object key -eq 'api_key').value\n"
    "$secret = ($e.values | Where-Object key -eq 'api_secret').value\n"
    "$auth = 'Basic ' + [Convert]::ToBase64String([Text.Encoding]::ASCII.GetBytes($key + ':' + $secret))",
)

doc.add_heading("Open the JMX in the JMeter GUI with runtime credentials", level=2)
add_code(
    doc,
    "& 'C:\\Tools\\apache-jmeter-5.6.3\\bin\\jmeter.bat' `\n"
    "  \"-JAuthorization=$auth\" `\n"
    "  -t 'C:\\QA\\output\\correlated.jmx'",
)
doc.add_paragraph("For a non-GUI validation run:")
add_code(
    doc,
    "& 'C:\\Tools\\apache-jmeter-5.6.3\\bin\\jmeter.bat' `\n"
    "  -n -t 'C:\\QA\\output\\correlated.jmx' `\n"
    "  \"-JAuthorization=$auth\" `\n"
    "  -l 'C:\\QA\\output\\results.jtl'",
)

doc.add_heading("Validation acceptance criteria", level=2)
for item in [
    "All required samplers execute.",
    "Authentication succeeds and there are no HTTP 401 or 403 responses.",
    "The producer response is successful.",
    "The extractor captures a value instead of __NOT_FOUND__.",
    "Downstream requests contain the extracted runtime value.",
    "The sampler error ratio meets the project threshold.",
]:
    add_bullet(doc, item)

doc.add_page_break()

# Handoff and security
doc.add_heading("7 Package the Deliverables", level=1)
doc.add_paragraph(
    "Give testers the files needed to reproduce the workflow while separating reusable assets from secrets."
)
add_table(
    doc,
    ["File", "Share", "Purpose"],
    [
        ("collection.postman_collection.json", "Yes", "Defines the API workflow."),
        ("environment.template.json", "Yes", "Documents required variables with empty values."),
        ("environment.json with real credentials", "No", "Contains secrets."),
        ("baseline.json and comparison.json", "Restricted", "Contain executed requests, responses and possibly credentials or personal data."),
        ("correlated.jmx", "Yes when secrets are external", "Runs the correlated JMeter workflow."),
        ("manifest.json", "Yes", "Explains variables, extractors, consumers and required properties."),
        ("run-jmeter.ps1 template", "Yes", "Provides a repeatable runtime command without embedded secret values."),
    ],
    [2.55, 1.0, 2.6],
)

doc.add_heading("Suggested tester package", level=2)
add_code(
    doc,
    "tester-package/\n"
    "  collection.postman_collection.json\n"
    "  environment.template.json\n"
    "  correlated.jmx\n"
    "  manifest.json\n"
    "  run-jmeter.ps1\n"
    "  Baseline11_Auto_Correlation_Tester_Guide.docx",
)

doc.add_heading("Security rules", level=2)
for item in [
    "Never commit filled environment files, Authorization values or API secrets to source control.",
    "Treat Newman reports as sensitive until request headers and bodies have been reviewed.",
    "Keep Embed static secrets disabled for shared or production test plans.",
    "Use test-mode credentials for demonstrations.",
    "Rotate a key if it has been exposed in a screenshot, report, ticket or repository.",
]:
    add_bullet(doc, item)

doc.add_page_break()

# Troubleshooting
doc.add_heading("8 Troubleshooting", level=1)
add_table(
    doc,
    ["Symptom", "Likely cause", "Corrective action"],
    [
        ("401 Unauthorized in Newman", "Environment values are empty, incorrect, or overridden by request-level auth.", "Verify exact variable names and request-level Basic Auth settings, then export the environment again."),
        ("URL contains null", "The producer failed or its Postman script did not set the variable.", "Fix the producer response and assertion before creating reports."),
        ("Literal {order_id} remains", "The collection uses API-documentation braces instead of Postman syntax.", "Use {{order_id}} in the Postman collection."),
        ("No candidates", "Values did not change, were not reused, or runs are misaligned.", "Use two fresh successful runs and verify the downstream request contains the produced value."),
        ("Duplicate ID candidates", "A later response echoes a value already produced earlier.", "Use a current Baseline11 build; the earliest original producer should be retained."),
        ("Ready with review", "A business heuristic detected an empty or error-like payload behind HTTP 2xx.", "Review the response. Continue when the empty result is valid for the scenario."),
        ("401 in downloaded JMX", "The external Authorization JMeter property was not supplied.", "Launch JMeter with -JAuthorization or provide the value in the validation form."),
        ("Extractor returns __NOT_FOUND__", "The response path or response content differs at runtime.", "Inspect the producer response and verify the extractor expression and match number."),
        ("Generate shows 0", "No candidate has been accepted.", "Run automatic correlation or accept the reviewed candidate."),
    ],
    [1.6, 2.15, 2.4],
)

doc.add_heading("Tester completion checklist", level=2)
checklist = [
    "[ ] Collection completes successfully in Newman.",
    "[ ] Baseline and comparison JSON reports were generated from the same scenario.",
    "[ ] Run Health has no blocker.",
    "[ ] Candidate producer and consumers match the intended business flow.",
    "[ ] Automatic correlation completed once.",
    "[ ] Dependency graph shows the expected edges.",
    "[ ] Generated JMX passes structural validation.",
    "[ ] JMeter execution succeeds with runtime properties.",
    "[ ] Extracted variables have real values and are not __NOT_FOUND__.",
    "[ ] No real credentials are included in the shared package.",
]
for item in checklist:
    add_bullet(doc, item)

doc.add_heading("Glossary", level=2)
add_table(
    doc,
    ["Term", "Meaning"],
    [
        ("Producer", "An earlier response that returns a dynamic value."),
        ("Consumer", "A later request that reuses the produced value."),
        ("Correlation", "Capturing a dynamic response value and substituting it into later requests."),
        ("Extractor", "A JMeter post-processor that reads a value from a response."),
        ("Parameterization", "Supplying user or environment input that has no response producer."),
        ("JMeter property", "A runtime value accessed with ${__P(name,)} and normally passed with -Jname=value."),
        ("Generated JMX", "A structurally valid plan that has not necessarily executed successfully."),
        ("Validated JMX", "A generated plan that Apache JMeter executed within the configured success threshold."),
    ],
    [1.65, 4.5],
)

doc.add_paragraph()
closing = doc.add_paragraph()
closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = closing.add_run("End of tester guide")
r.bold = True
r.font.color.rgb = GRAY

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
